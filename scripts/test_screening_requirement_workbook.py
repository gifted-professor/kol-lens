#!/usr/bin/env python3
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from screening_requirement_workbook import (
    MAIN_SHEET_NAME,
    PREVIEW_SHEET_NAME,
    OPTION_SHEET_NAME,
    compile_requirement_workbook,
    create_legacy_template_workbook,
    parse_requirement_workbook,
    save_template,
)


def fill_sectioned_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook[MAIN_SHEET_NAME]

    sheet["B5"] = "Tapo 北美生活方式达人筛号"
    sheet["B6"] = "Tapo 智能家居"
    sheet["B7"] = "两者"
    sheet["B8"] = "判断达人是否符合家庭 / 宠物 / 户外生活场景合作标准"
    sheet["B9"] = "creator_a,creator_b"
    sheet["B10"] = "creator_x"
    sheet["B11"] = "测试用示例"

    sheet["B15"] = "美国,加拿大"
    sheet["B16"] = "英语,English"
    sheet["B17"] = "主页资料"
    sheet["B18"] = "直接不通过"

    sheet["B23"] = "50"
    sheet["B24"] = "10000"
    sheet["B25"] = "10000"
    sheet["B28"] = "同时满足"
    sheet["B29"] = "直接不通过"

    sheet["B34"] = "18"
    sheet["B35"] = "1"
    sheet["B38"] = "是"
    sheet["C38"] = "街访、朋友互动"
    sheet["B39"] = "是"
    sheet["C39"] = "镜头前开口说话"
    sheet["B42"] = "是"
    sheet["C42"] = "手持或展示产品"
    sheet["B43"] = "是"
    sheet["C43"] = "户外庭院场景"
    sheet["B44"] = "是"
    sheet["C44"] = "与宠物互动"

    sheet["B48"] = "beauty,makeup,skincare"
    sheet["B49"] = "partner,boyfriend,girlfriend,husband,wife"
    sheet["B50"] = "50"
    sheet["B51"] = "30"
    sheet["B52"] = "70"
    sheet["B53"] = "是"
    sheet["B54"] = "solo POV 内容占比过高时人工复核"

    sheet["B59"] = "是"
    sheet["B60"] = "是"
    sheet["B61"] = "是"
    sheet["B62"] = "solo POV 是否缺少互动感"
    sheet["B63"] = "若需求涉及年龄或种族，只保留合规提醒"

    sheet["B68"] = "步骤1：基础资质审核,步骤2：数据审核,步骤3：内容 / 视觉审核"
    sheet["B69"] = "步骤4：排除项审核"
    sheet["B70"] = "转人工"
    sheet["B71"] = "通过"
    sheet["B72"] = "不通过"

    workbook.save(path)


def test_build_outputs():
    with TemporaryDirectory() as temp_dir:
        base = Path(temp_dir)
        template_path = base / "screening-template.xlsx"
        spec_path = base / "screening-template.md"

        result = save_template(template_path, spec_path)

        assert Path(result["template_path"]).exists()
        assert Path(result["spec_path"]).exists()
        assert Path(result["contract_path"]).exists()

        workbook = load_workbook(template_path)
        assert MAIN_SHEET_NAME in workbook.sheetnames
        assert PREVIEW_SHEET_NAME in workbook.sheetnames
        assert OPTION_SHEET_NAME in workbook.sheetnames
        assert workbook[PREVIEW_SHEET_NAME].sheet_state == "hidden"

        spec_text = spec_path.read_text(encoding="utf-8")
        assert "填写手册" in spec_text
        assert "先记住 3 件事" in spec_text
        assert "不会填时怎么判断" in spec_text


def test_parse_sectioned_workbook():
    with TemporaryDirectory() as temp_dir:
        base = Path(temp_dir)
        template_path = base / "screening-template.xlsx"
        spec_path = base / "screening-template.md"
        save_template(template_path, spec_path)
        fill_sectioned_workbook(template_path)

        parsed = parse_requirement_workbook(template_path)
        assert parsed["layout"] == "sectioned_main_sheet_v1"
        assert not parsed["errors"]

        structured = parsed["structured_requirement"]
        assert structured["basic_info"]["project_name"] == "Tapo 北美生活方式达人筛号"
        assert structured["basic_info"]["platform_scope"] == ["tiktok", "instagram"]
        assert structured["qualification"]["region_requirement"] == ["美国", "加拿大"]
        assert structured["data_audit"]["mean_view_threshold"] == 10000
        assert structured["visual_audit"]["features"][0]["enabled"] is True
        assert structured["manual_review"]["persona_or_niche_manual"] is True
        assert structured["final_logic"]["success_output"] == "pass"


def test_compile_sectioned_workbook():
    with TemporaryDirectory() as temp_dir:
        base = Path(temp_dir)
        template_path = base / "screening-template.xlsx"
        spec_path = base / "screening-template.md"
        output_dir = base / "compiled"
        save_template(template_path, spec_path)
        fill_sectioned_workbook(template_path)

        compiled = compile_requirement_workbook(template_path, output_dir=str(output_dir))
        assert compiled["success"] is True
        assert not compiled["errors"]
        assert Path(compiled["artifacts"]["structured_requirement_json"]).exists()
        assert Path(compiled["artifacts"]["rulespec_json"]).exists()
        assert Path(compiled["artifacts"]["compiled_workbook"]).exists()

        rulespec = compiled["rulespec"]
        rule_types = {rule["rule_type"] for rule in rulespec["rules"]}
        assert "visual_feature_group" in rule_types
        assert "region_gate" in rule_types
        assert "language_gate" in rule_types
        assert "content_keyword_ratio" in rule_types
        assert all(note["key"] == "protected_attribute_notice" for note in rulespec["compliance_notes"])
        assert all(rule["rule_type"] != "protected_attribute_notice" for rule in rulespec["rules"])

        workbook = load_workbook(compiled["artifacts"]["compiled_workbook"])
        preview_sheet = workbook[PREVIEW_SHEET_NAME]
        assert preview_sheet.sheet_state == "hidden"
        preview_values = [preview_sheet[f"E{row}"].value for row in range(2, preview_sheet.max_row + 1)]
        assert "compiled" in preview_values
        assert "manual_only" in preview_values
        assert "compliance_only" in preview_values


def test_parse_legacy_workbook_rows():
    with TemporaryDirectory() as temp_dir:
        base = Path(temp_dir)
        template_path = base / "legacy-template.xlsx"

        workbook = create_legacy_template_workbook()
        auto_sheet = workbook["自动初筛规则"]
        auto_sheet["A2"] = "R001"
        auto_sheet["B2"] = "是"
        auto_sheet["C2"] = "TikTok"
        auto_sheet["D2"] = "基础数据"
        auto_sheet["E2"] = "步骤1 基础数据"
        auto_sheet["F2"] = "播放量中位数"
        auto_sheet["G2"] = "最近内容"
        auto_sheet["H2"] = "50"
        auto_sheet["I2"] = "条"
        auto_sheet["J2"] = "大于"
        auto_sheet["K2"] = "10000"
        auto_sheet["L2"] = "次"
        auto_sheet["N2"] = "不通过"
        workbook.save(template_path)

        parsed = parse_requirement_workbook(template_path)
        assert parsed["layout"] == "legacy_multi_sheet"
        assert not parsed["errors"]
        assert parsed["sheets"]["automatic_rules"][0]["normalized"]["judgement_item"] == "view_count_median"


if __name__ == "__main__":
    test_build_outputs()
    test_parse_sectioned_workbook()
    test_compile_sectioned_workbook()
    test_parse_legacy_workbook_rows()
    print("ok")
