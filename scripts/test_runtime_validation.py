import json
import os
import sys
import tempfile
import threading
import time
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import backend.app as backend_app


@contextmanager
def patched_runtime_env():
    original_auth_file = backend_app.APIFY_AUTH_FILE
    original_data_dir = backend_app.DATA_DIR
    original_vision_provider_state_file = backend_app.VISION_PROVIDER_STATE_FILE
    original_upload_folder = backend_app.UPLOAD_FOLDER
    original_app_upload_folder = backend_app.app.config.get("UPLOAD_FOLDER")
    scripts_dir = REPO_ROOT / "scripts"
    if str(scripts_dir) not in sys.path:
        sys.path.insert(0, str(scripts_dir))
    try:
        import data_cleaner as screening_data_cleaner
    except Exception:
        screening_data_cleaner = None
    original_cleaner_data_dir = getattr(screening_data_cleaner, "DATA_DIR", None) if screening_data_cleaner else None
    env_keys = [
        "APIFY_TOKEN",
        "APIFY_API_TOKEN",
        "APIFY_BACKUP_TOKENS",
        "APIFY_TOKENS",
        "OPENAI_API_KEY",
        "VISION_LEMONAPI_API_KEY",
        "VISION_LEMONAPI_BASE_URL",
        "VISION_LEMONAPI_MODEL",
        "VISION_QUAN2GO_API_KEY",
    ]
    original_env = {key: os.environ.get(key) for key in env_keys}

    with tempfile.TemporaryDirectory() as temp_dir:
        missing_auth_file = Path(temp_dir) / "missing-auth.json"
        temp_data_dir = Path(temp_dir) / "data"
        temp_upload_dir = temp_data_dir / "uploads"
        temp_upload_dir.mkdir(parents=True, exist_ok=True)
        backend_app.APIFY_AUTH_FILE = str(missing_auth_file)
        backend_app.DATA_DIR = str(temp_data_dir)
        backend_app.VISION_PROVIDER_STATE_FILE = str(temp_data_dir / "vision_provider_state.json")
        backend_app.UPLOAD_FOLDER = str(temp_upload_dir)
        backend_app.app.config["UPLOAD_FOLDER"] = str(temp_upload_dir)
        if screening_data_cleaner is not None:
            screening_data_cleaner.DATA_DIR = str(temp_data_dir)
        for key in env_keys:
            os.environ.pop(key, None)
        try:
            yield
        finally:
            backend_app.APIFY_AUTH_FILE = original_auth_file
            backend_app.DATA_DIR = original_data_dir
            backend_app.VISION_PROVIDER_STATE_FILE = original_vision_provider_state_file
            backend_app.UPLOAD_FOLDER = original_upload_folder
            backend_app.app.config["UPLOAD_FOLDER"] = original_app_upload_folder
            if screening_data_cleaner is not None and original_cleaner_data_dir is not None:
                screening_data_cleaner.DATA_DIR = original_cleaner_data_dir
            for key, value in original_env.items():
                if value is None:
                    os.environ.pop(key, None)
                else:
                    os.environ[key] = value


def assert_json_error(response, expected_error_code, expected_substring):
    payload = response.get_json()
    assert response.status_code == 400, payload
    assert payload["success"] is False, payload
    assert payload["error_code"] == expected_error_code, payload
    assert expected_substring in payload["error"], payload
    assert isinstance(payload.get("details"), list) and payload["details"], payload


def post_upload_workbook(client, workbook_path):
    with open(workbook_path, "rb") as f:
        response = client.post(
            "/api/upload",
            data={"file": (f, workbook_path.name)},
            content_type="multipart/form-data",
        )
    return response


def build_missing_platform_workbook(source_path):
    temp_dir = Path(tempfile.mkdtemp())
    temp_path = temp_dir / "missing_platform.xlsx"
    workbook = load_workbook(source_path)
    sheet = workbook[workbook.sheetnames[0]]
    for cell in sheet[1]:
        if cell.value == "Platform":
            cell.value = "Channel"
            break
    workbook.save(temp_path)
    return temp_path


def build_multisheet_canonical_workbook():
    temp_dir = Path(tempfile.mkdtemp())
    temp_path = temp_dir / "multi_sheet_upload.xlsx"

    workbook = Workbook()
    headers = [
        "Platform",
        "@username",
        "nickname",
        "Region",
        "Language",
        "Followers",
        "URL",
    ]
    sheet_rows = {
        "YouTube": [["YouTube", "@creatoralpha", "Alpha", "US", "en", 100000, ""]],
        "TikTok": [["TikTok", "@creatorbeta", "Beta", "CA", "en", 200000, ""]],
        "Instagram": [["Instagram", "@creatorgamma", "Gamma", "US", "en", 300000, ""]],
    }

    first_sheet = workbook.active
    first_sheet.title = "YouTube"
    created_sheets = {"YouTube": first_sheet}

    for sheet_name in ("TikTok", "Instagram"):
        created_sheets[sheet_name] = workbook.create_sheet(title=sheet_name)

    for sheet_name, rows in sheet_rows.items():
        sheet = created_sheets[sheet_name]
        sheet.append(headers)
        for row in rows:
            sheet.append(row)

    workbook.save(temp_path)
    return temp_path


def write_json(path, payload):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def workbook_sheet_headers(workbook_bytes, sheet_name=None):
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    sheet = workbook[sheet_name or workbook.sheetnames[0]]
    return [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]


def workbook_sheet_row_dict(workbook_bytes, sheet_name=None, row_index=2):
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    sheet = workbook[sheet_name or workbook.sheetnames[0]]
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    values = [sheet.cell(row_index, column).value for column in range(1, sheet.max_column + 1)]
    return dict(zip(headers, values))


def wait_for(predicate, timeout=2.0, interval=0.02):
    started_at = time.time()
    while time.time() - started_at < timeout:
        value = predicate()
        if value:
            return value
        time.sleep(interval)
    raise AssertionError("Timed out waiting for condition")


@contextmanager
def patched_backend_attrs(**patches):
    originals = {name: getattr(backend_app, name) for name in patches}
    try:
        for name, value in patches.items():
            setattr(backend_app, name, value)
        yield
    finally:
        for name, value in originals.items():
            setattr(backend_app, name, value)


def main():
    with patched_runtime_env():
        import data_cleaner

        client = backend_app.app.test_client()
        sample_workbook = REPO_ROOT / "data" / "uploads" / "kol_emails_2026-03-16-07-08-05_xz406658gmail.com.xlsx"

        upload_response = post_upload_workbook(client, sample_workbook)
        upload_payload = upload_response.get_json()
        assert upload_response.status_code == 200, upload_payload
        assert upload_payload["success"] is True, upload_payload
        assert upload_payload["stats"]["Instagram"] > 0, upload_payload
        assert upload_payload["stats"]["TikTok"] == 0, upload_payload
        assert upload_payload["metadata_counts"]["instagram"] == upload_payload["stats"]["Instagram"], upload_payload
        assert upload_payload["grouped_data"]["instagram"], upload_payload
        assert upload_payload["grouped_data"]["instagram"][0].startswith("https://instagram.com/"), upload_payload

        instagram_metadata_path = Path(backend_app.get_upload_metadata_path("instagram"))
        assert instagram_metadata_path.exists(), instagram_metadata_path
        instagram_metadata = backend_app.load_upload_metadata("instagram")
        assert instagram_metadata["lifewithhaaliyah"]["region"] == "CA", instagram_metadata["lifewithhaaliyah"]
        assert instagram_metadata["lifewithhaaliyah"]["language"] == "en", instagram_metadata["lifewithhaaliyah"]
        assert instagram_metadata["lifewithhaaliyah"]["avg_views"] == 298264, instagram_metadata["lifewithhaaliyah"]

        multisheet_workbook = build_multisheet_canonical_workbook()
        multisheet_upload_response = post_upload_workbook(client, multisheet_workbook)
        multisheet_payload = multisheet_upload_response.get_json()
        assert multisheet_upload_response.status_code == 200, multisheet_payload
        assert multisheet_payload["success"] is True, multisheet_payload
        assert multisheet_payload["stats"]["YouTube"] == 1, multisheet_payload
        assert multisheet_payload["stats"]["TikTok"] == 1, multisheet_payload
        assert multisheet_payload["stats"]["Instagram"] == 1, multisheet_payload
        assert multisheet_payload["metadata_counts"]["youtube"] == 1, multisheet_payload
        assert multisheet_payload["metadata_counts"]["tiktok"] == 1, multisheet_payload
        assert multisheet_payload["metadata_counts"]["instagram"] == 1, multisheet_payload
        assert multisheet_payload["grouped_data"]["youtube"] == ["https://www.youtube.com/@creatoralpha"], multisheet_payload
        assert multisheet_payload["grouped_data"]["tiktok"] == ["https://www.tiktok.com/@creatorbeta"], multisheet_payload
        assert multisheet_payload["grouped_data"]["instagram"] == ["https://www.instagram.com/creatorgamma/"], multisheet_payload

        youtube_metadata = backend_app.load_upload_metadata("youtube")
        tiktok_metadata = backend_app.load_upload_metadata("tiktok")
        instagram_metadata = backend_app.load_upload_metadata("instagram")
        assert youtube_metadata["creatoralpha"]["region"] == "US", youtube_metadata["creatoralpha"]
        assert tiktok_metadata["creatorbeta"]["region"] == "CA", tiktok_metadata["creatorbeta"]
        assert instagram_metadata["creatorgamma"]["region"] == "US", instagram_metadata["creatorgamma"]

        invalid_workbook = build_missing_platform_workbook(sample_workbook)
        invalid_upload_response = post_upload_workbook(client, invalid_workbook)
        invalid_payload = invalid_upload_response.get_json()
        assert invalid_upload_response.status_code == 400, invalid_payload
        assert invalid_payload["error_code"] == "UPLOAD_TEMPLATE_INVALID", invalid_payload
        assert "Platform" in invalid_payload["error"], invalid_payload
        assert any("缺少必填列：Platform" in detail for detail in invalid_payload["details"]), invalid_payload

        expected_review_keys = {
            "platform",
            "username",
            "profile_url",
            "status",
            "reason",
            "covers",
            "latest_post_time",
            "soft_flags",
            "stats",
            "upload_metadata",
        }

        backend_app.save_upload_metadata(
            "instagram",
            {
                "creatorgamma": {
                    "handle": "creatorgamma",
                    "url": "https://www.instagram.com/creatorgamma",
                    "region": "US",
                    "language": "en",
                    "source_filename": "runtime_validation_upload.xlsx",
                }
            },
        )

        tiktok_contract_path = Path(backend_app.get_platform_dir("tiktok")) / "tiktok_contract.json"
        write_json(
            tiktok_contract_path,
            [
                {
                    "authorMeta": {
                        "name": "creatorbeta",
                        "profileUrl": "https://www.tiktok.com/@creatorbeta",
                    },
                    "createTimeISO": "2026-03-16T00:00:00Z",
                    "playCount": 25000,
                    "text": "smart home setup",
                    "hashtags": [{"name": "home"}],
                    "videoMeta": {"coverUrl": "https://example.com/tiktok-cover.jpg"},
                }
            ],
        )
        tiktok_contract_result = data_cleaner.filter_and_save_dataset(
            str(tiktok_contract_path),
            "tiktok",
            ["https://www.tiktok.com/@creatorbeta", "https://www.tiktok.com/@missingbeta"],
        )
        assert tiktok_contract_result["success"] is True, tiktok_contract_result
        assert {item["status"] for item in tiktok_contract_result["profile_reviews"]} == {"Pass", "Missing"}, tiktok_contract_result
        for item in tiktok_contract_result["profile_reviews"]:
            assert expected_review_keys.issubset(item.keys()), item
            assert item["platform"] == "tiktok", item
            assert isinstance(item["covers"], list), item
            assert isinstance(item["soft_flags"], list), item
            assert isinstance(item["stats"], dict), item
            assert isinstance(item["upload_metadata"], dict), item

        instagram_contract_path = Path(backend_app.get_platform_dir("instagram")) / "instagram_contract.json"
        write_json(
            instagram_contract_path,
            [
                {
                    "username": "creatorgamma",
                    "url": "https://www.instagram.com/creatorgamma",
                    "biography": "",
                    "latestPosts": [
                        {
                            "timestamp": "2026-03-16T00:00:00Z",
                            "caption": "daily routine",
                            "displayUrl": "https://example.com/ig-cover.jpg",
                        }
                    ],
                }
            ],
        )
        instagram_contract_result = data_cleaner.filter_and_save_dataset(
            str(instagram_contract_path),
            "instagram",
            ["https://www.instagram.com/creatorgamma", "https://www.instagram.com/missinggamma"],
        )
        assert instagram_contract_result["success"] is True, instagram_contract_result
        assert {item["status"] for item in instagram_contract_result["profile_reviews"]} == {"Pass", "Missing"}, instagram_contract_result
        for item in instagram_contract_result["profile_reviews"]:
            assert expected_review_keys.issubset(item.keys()), item
            assert item["platform"] == "instagram", item
            assert isinstance(item["covers"], list), item
            assert isinstance(item["soft_flags"], list), item
            assert isinstance(item["stats"], dict), item
            assert isinstance(item["upload_metadata"], dict), item

        youtube_contract_path = Path(backend_app.get_platform_dir("youtube")) / "youtube_contract.json"
        write_json(
            youtube_contract_path,
            [
                {
                    "channelName": "creatoralpha",
                    "channelUrl": "https://www.youtube.com/@creatoralpha",
                    "date": "2026-03-16T00:00:00Z",
                    "title": "smart home vlog",
                    "text": "clean content",
                    "thumbnailUrl": "https://example.com/yt-cover.jpg",
                    "descriptionLinks": [],
                    "isPaidContent": False,
                    "aboutChannelInfo": {
                        "channelDescription": "tech creator",
                        "channelDescriptionLinks": [],
                    },
                }
            ],
        )
        youtube_contract_result = data_cleaner.filter_and_save_dataset(
            str(youtube_contract_path),
            "youtube",
            ["https://www.youtube.com/@creatoralpha", "https://www.youtube.com/@missingalpha"],
        )
        assert youtube_contract_result["success"] is True, youtube_contract_result
        assert {item["status"] for item in youtube_contract_result["profile_reviews"]} == {"Pass", "Missing"}, youtube_contract_result
        for item in youtube_contract_result["profile_reviews"]:
            assert expected_review_keys.issubset(item.keys()), item
            assert item["platform"] == "youtube", item
            assert isinstance(item["covers"], list), item
            assert isinstance(item["soft_flags"], list), item
            assert isinstance(item["stats"], dict), item
            assert isinstance(item["upload_metadata"], dict), item

        fake_job = backend_app.create_job("scrape", platform="tiktok", message="采集任务已创建")
        runner_started = threading.Event()

        def fake_scrape_runner(progress_callback, cancel_check):
            runner_started.set()
            progress_callback(
                "batch_completed",
                "已完成第 1/2 批",
                done=1,
                total=2,
                batch_index=1,
                batch_total=2,
                partial_result={
                    "success": True,
                    "is_partial": True,
                    "count": 1,
                    "profile_reviews": [],
                    "requested_total": 2,
                    "completed_batches": 1,
                    "total_batches": 2,
                    "failed_batches": [{"batch_index": 2, "error": "demo failure"}],
                    "message": "partial",
                },
            )
            while not cancel_check():
                time.sleep(0.01)
            return backend_app.build_cancelled_result("用户取消")

        backend_app.start_background_job(fake_job, fake_scrape_runner)

        wait_for(lambda: runner_started.is_set())
        running_job = wait_for(lambda: backend_app.get_job(fake_job["id"]).get("partial_result"))
        latest_running_job = backend_app.get_job(fake_job["id"])
        assert latest_running_job["status"] == "running", latest_running_job
        assert latest_running_job["contract_version"] == "scrape_job_v1", latest_running_job
        assert latest_running_job["progress"]["determinate"] is True, latest_running_job
        assert latest_running_job["batch_summary"]["completed"] == 1, latest_running_job
        assert latest_running_job["failed_batches"][0]["error"] == "demo failure", latest_running_job

        cancel_response = client.post(f"/api/jobs/{fake_job['id']}/cancel")
        cancel_payload = cancel_response.get_json()
        assert cancel_response.status_code == 200, cancel_payload
        assert cancel_payload["job"]["status"] == "cancelling", cancel_payload
        assert cancel_payload["job"]["stage"] == "cancelling", cancel_payload
        assert cancel_payload["job"]["cancel_requested_at"], cancel_payload

        cancelled_job = wait_for(lambda: backend_app.get_job(fake_job["id"]) if backend_app.get_job(fake_job["id"])["status"] == "cancelled" else None)
        assert cancelled_job["stage"] == "cancelled", cancelled_job
        assert cancelled_job["partial_result"]["is_partial"] is True, cancelled_job
        assert cancelled_job["failed_batches"][0]["error"] == "demo failure", cancelled_job
        assert cancelled_job["finished_at"], cancelled_job

        youtube_dir = Path(backend_app.get_platform_dir("youtube"))
        raw_current = youtube_dir / "youtube_data.json"
        raw_snapshot = youtube_dir / "youtube_data_last_non_empty.json"
        profile_reviews_path = youtube_dir / "youtube_profile_reviews.json"
        raw_current.write_text("[]", encoding="utf-8")
        raw_snapshot.write_text('[{"channelName":"demo-channel","id":"abc"}]', encoding="utf-8")
        profile_reviews_path.write_text(
            '[{"username":"demo-channel","status":"Pass","reason":"","covers":[]}]',
            encoding="utf-8",
        )

        loaded_youtube_reviews = backend_app.load_profile_reviews("youtube")
        assert loaded_youtube_reviews[0]["platform"] == "youtube", loaded_youtube_reviews
        assert isinstance(loaded_youtube_reviews[0]["soft_flags"], list), loaded_youtube_reviews
        assert isinstance(loaded_youtube_reviews[0]["stats"], dict), loaded_youtube_reviews
        assert isinstance(loaded_youtube_reviews[0]["upload_metadata"], dict), loaded_youtube_reviews

        cached_snapshot_result = backend_app.build_cached_scrape_response("youtube", str(raw_current), 3)
        assert cached_snapshot_result["success"] is True, cached_snapshot_result
        assert cached_snapshot_result["cached"] is True, cached_snapshot_result
        assert cached_snapshot_result["used_fallback"] is True, cached_snapshot_result
        assert cached_snapshot_result["raw_data_source"] == "last_non_empty_snapshot", cached_snapshot_result
        assert "最近一次非空快照结果" in cached_snapshot_result["message"], cached_snapshot_result

        preserved_failure_result = backend_app.build_preserved_failure_response(
            "youtube",
            str(raw_current),
            [{"batch_index": 1, "error": "network fail"}],
        )
        assert preserved_failure_result["success"] is True, preserved_failure_result
        assert preserved_failure_result["stale_result"] is True, preserved_failure_result
        assert preserved_failure_result["used_fallback"] is True, preserved_failure_result
        assert preserved_failure_result["failed_batches"][0]["error"] == "network fail", preserved_failure_result
        assert "已保留最近一次可用结果" in preserved_failure_result["message"], preserved_failure_result

        instagram_raw_path = Path(backend_app.get_platform_dir("instagram")) / "instagram_data.json"
        instagram_raw_snapshot_path = Path(backend_app.get_last_non_empty_raw_snapshot_path("instagram"))

        cached_subset_reviews = [
            {
                "platform": "instagram",
                "username": "cachedalpha",
                "profile_url": "https://www.instagram.com/cachedalpha",
                "status": "Pass",
                "reason": "cached alpha",
                "covers": ["https://example.com/cached-alpha.jpg"],
                "latest_post_time": "2026-03-16T00:00:00Z",
                "soft_flags": [],
                "stats": {},
                "upload_metadata": {},
            },
            {
                "platform": "instagram",
                "username": "cachedbeta",
                "profile_url": "https://www.instagram.com/cachedbeta",
                "status": "Pass",
                "reason": "cached beta",
                "covers": ["https://example.com/cached-beta.jpg"],
                "latest_post_time": "2026-03-16T00:00:00Z",
                "soft_flags": [],
                "stats": {},
                "upload_metadata": {},
            },
            {
                "platform": "instagram",
                "username": "cachedextra",
                "profile_url": "https://www.instagram.com/cachedextra",
                "status": "Pass",
                "reason": "cached extra",
                "covers": ["https://example.com/cached-extra.jpg"],
                "latest_post_time": "2026-03-16T00:00:00Z",
                "soft_flags": [],
                "stats": {},
                "upload_metadata": {},
            },
        ]
        write_json(
            instagram_raw_path,
            [
                {
                    "username": "cachedalpha",
                    "url": "https://www.instagram.com/cachedalpha",
                    "biography": "",
                    "latestPosts": [{"timestamp": "2026-03-16T00:00:00Z", "caption": "alpha", "displayUrl": "https://example.com/cached-alpha.jpg"}],
                },
                {
                    "username": "cachedbeta",
                    "url": "https://www.instagram.com/cachedbeta",
                    "biography": "",
                    "latestPosts": [{"timestamp": "2026-03-16T00:00:00Z", "caption": "beta", "displayUrl": "https://example.com/cached-beta.jpg"}],
                },
                {
                    "username": "cachedextra",
                    "url": "https://www.instagram.com/cachedextra",
                    "biography": "",
                    "latestPosts": [{"timestamp": "2026-03-16T00:00:00Z", "caption": "extra", "displayUrl": "https://example.com/cached-extra.jpg"}],
                },
            ],
        )
        backend_app.save_profile_reviews("instagram", cached_subset_reviews)

        cached_subset_result = backend_app.build_cached_scrape_response(
            "instagram",
            str(instagram_raw_path),
            2,
            requested_identifiers=[
                "https://www.instagram.com/cachedbeta",
                "https://www.instagram.com/cachedalpha",
            ],
        )
        assert cached_subset_result["success"] is True, cached_subset_result
        assert cached_subset_result["requested_total"] == 2, cached_subset_result
        assert [item["username"] for item in cached_subset_result["profile_reviews"]] == [
            "cachedbeta",
            "cachedalpha",
        ], cached_subset_result

        backend_app.save_profile_reviews("instagram", cached_subset_reviews[:2])
        backend_app.save_upload_metadata(
            "instagram",
            {
                "newfresh": {
                    "handle": "newfresh",
                    "url": "https://www.instagram.com/newfresh",
                    "region": "US",
                    "language": "en",
                    "source_filename": "runtime_validation_upload.xlsx",
                }
            },
        )
        write_json(
            instagram_raw_path,
            [
                {
                    "username": "newfresh",
                    "url": "https://www.instagram.com/newfresh",
                    "biography": "",
                    "latestPosts": [
                        {
                            "timestamp": "2026-03-16T00:00:00Z",
                            "caption": "new fresh",
                            "displayUrl": "https://example.com/newfresh-cover.jpg",
                        }
                    ],
                }
            ],
        )

        mixed_cache_result = backend_app.finalize_apify_output(
            "instagram",
            str(instagram_raw_path),
            ["https://www.instagram.com/newfresh"],
            2,
            requested_identifiers=[
                "https://www.instagram.com/cachedalpha",
                "https://www.instagram.com/cachedbeta",
                "https://www.instagram.com/newfresh",
            ],
        )
        assert mixed_cache_result["success"] is True, mixed_cache_result
        assert mixed_cache_result["requested_total"] == 3, mixed_cache_result
        assert [item["username"] for item in mixed_cache_result["profile_reviews"]] == [
            "cachedalpha",
            "cachedbeta",
            "newfresh",
        ], mixed_cache_result
        assert len(mixed_cache_result["successful_identifiers"]) == 3, mixed_cache_result
        assert all(item["status"] == "Pass" for item in mixed_cache_result["profile_reviews"]), mixed_cache_result

        stale_instagram_reviews = [
            {
                "platform": "instagram",
                "username": "creatorgamma",
                "profile_url": "",
                "status": "Pass",
                "reason": "",
                "covers": ["https://example.com/ig-cover.jpg"],
                "latest_post_time": "2026-03-16T00:00:00Z",
                "soft_flags": [],
                "stats": {},
                "upload_metadata": {
                    "handle": "creatorgamma",
                    "url": "https://www.instagram.com/wronggamma",
                    "region": "FR",
                    "language": "fr",
                },
            }
        ]
        backend_app.save_profile_reviews(
            "instagram",
            stale_instagram_reviews,
            metadata={
                "cached": True,
                "stale_result": True,
                "used_fallback": True,
                "raw_data_source": "last_non_empty_snapshot",
                "source_path": str(instagram_raw_snapshot_path),
            },
        )

        canonical_instagram_reviews = backend_app.load_profile_reviews("instagram")
        assert canonical_instagram_reviews[0]["upload_metadata"]["region"] == "US", canonical_instagram_reviews
        assert canonical_instagram_reviews[0]["upload_metadata"]["language"] == "en", canonical_instagram_reviews
        assert canonical_instagram_reviews[0]["profile_url"] == "https://www.instagram.com/creatorgamma", canonical_instagram_reviews

        prescreen_rows = backend_app.build_prescreen_review_rows("instagram", stale_instagram_reviews)
        image_rows = backend_app.build_image_review_rows("instagram", stale_instagram_reviews)
        summary_rows = backend_app.build_test_info_summary_rows("instagram", stale_instagram_reviews, [], instagram_metadata)
        raw_rows = backend_app.build_test_info_raw_rows(
            "instagram",
            [{"username": "creatorgamma", "url": "https://www.instagram.com/creatorgamma"}],
            stale_instagram_reviews,
            instagram_metadata,
        )
        json_payload = backend_app.build_test_info_json_payload(
            "instagram",
            stale_instagram_reviews,
            [],
            instagram_metadata,
        )
        final_rows = backend_app.build_final_review_rows(
            "instagram",
            stale_instagram_reviews,
            {
                "creatorgamma": {
                    "username": "creatorgamma",
                    "decision": "Pass",
                    "reason": "looks good",
                    "signals": [],
                }
            },
        )

        for row in (prescreen_rows[0], image_rows[0], summary_rows[0], raw_rows[0], final_rows[0]):
            assert row["upload_region"] == "US", row
            assert row["upload_language"] == "en", row
            assert row["upload_handle"] == "creatorgamma", row

        assert final_rows[0]["profile_url"] == "https://www.instagram.com/creatorgamma", final_rows[0]
        payload_profile = next(item for item in json_payload["profiles"] if item["identifier"] == "creatorgamma")
        assert payload_profile["upload_metadata"]["region"] == "US", payload_profile
        assert payload_profile["review"]["upload_metadata"]["region"] == "US", payload_profile
        assert payload_profile["profile_url"] == "https://www.instagram.com/creatorgamma", payload_profile

        final_review_response = client.post(
            "/api/download/instagram/final-review",
            json={
                "profile_reviews": [
                    {
                        "platform": "instagram",
                        "username": "creatorgamma",
                        "profile_url": "",
                        "status": "Pass",
                        "reason": "",
                        "covers": [],
                        "upload_metadata": {
                            "handle": "creatorgamma",
                            "region": "ZZ",
                            "language": "zz",
                        },
                    }
                ],
                "visual_results": {
                    "creatorgamma": {
                        "username": "creatorgamma",
                        "decision": "Pass",
                        "reason": "ok",
                        "signals": [],
                    }
                },
            },
        )
        assert final_review_response.status_code == 200, final_review_response.get_data(as_text=True)
        final_review_headers = workbook_sheet_headers(final_review_response.data)
        final_review_row = workbook_sheet_row_dict(final_review_response.data)
        assert final_review_row["upload_region"] == "US", final_review_row
        assert final_review_row["upload_language"] == "en", final_review_row
        assert final_review_row["upload_handle"] == "creatorgamma", final_review_row
        assert final_review_row["profile_url"] == "https://www.instagram.com/creatorgamma", final_review_row
        assert final_review_row["source_filename"] == "runtime_validation_upload.xlsx", final_review_row
        assert {"identifier", "profile_url", "source_filename", "status", "reason", "account_id"}.issubset(final_review_headers), final_review_headers
        assert "stale_result" in final_review_headers, final_review_headers
        assert "raw_data_source" in final_review_headers, final_review_headers

        raw_route_items = [
            {
                "username": "creatorgamma",
                "url": "https://www.instagram.com/creatorgamma",
                "biography": "raw route payload",
                "latestPosts": [
                    {
                        "timestamp": "2026-03-16T00:00:00Z",
                        "caption": "route test",
                        "displayUrl": "https://example.com/ig-cover.jpg",
                    }
                ],
            }
        ]
        write_json(instagram_raw_path, raw_route_items)
        write_json(instagram_raw_snapshot_path, [{"username": "snapshotonly", "url": "https://www.instagram.com/snapshotonly"}])

        # Raw exports are the existing download_results(..., format) routes and should stay raw.
        raw_json_response = client.get("/api/download/instagram/json")
        assert raw_json_response.status_code == 200, raw_json_response.get_data(as_text=True)
        assert json.loads(raw_json_response.get_data(as_text=True)) == raw_route_items, raw_json_response.get_data(as_text=True)

        raw_excel_response = client.get("/api/download/instagram/excel")
        assert raw_excel_response.status_code == 200, raw_excel_response.get_data(as_text=True)
        raw_excel_headers = workbook_sheet_headers(raw_excel_response.data)
        assert {"username", "url", "biography", "latestPosts"}.issubset(raw_excel_headers), raw_excel_headers
        assert "identifier" not in raw_excel_headers, raw_excel_headers
        assert "profile_url" not in raw_excel_headers, raw_excel_headers
        assert "source_filename" not in raw_excel_headers, raw_excel_headers

        # Raw routes remain raw but now fall back to the last non-empty snapshot when current raw data is empty.
        write_json(instagram_raw_path, [])
        write_json(instagram_raw_snapshot_path, raw_route_items)

        raw_json_empty_response = client.get("/api/download/instagram/json")
        assert raw_json_empty_response.status_code == 200, raw_json_empty_response.get_data(as_text=True)
        assert json.loads(raw_json_empty_response.get_data(as_text=True)) == raw_route_items, raw_json_empty_response.get_data(as_text=True)

        raw_excel_empty_response = client.get("/api/download/instagram/excel")
        assert raw_excel_empty_response.status_code == 200
        raw_excel_empty_headers = workbook_sheet_headers(raw_excel_empty_response.data)
        assert {"username", "url", "biography", "latestPosts"}.issubset(raw_excel_empty_headers), raw_excel_empty_headers

        # test-info-json remains the structured debug/test JSON export, distinct from the raw json route above.
        test_info_response = client.get("/api/download/instagram/test-info")
        assert test_info_response.status_code == 200, test_info_response.get_data(as_text=True)
        test_info_account_headers = workbook_sheet_headers(test_info_response.data, "Account Review")
        test_info_account_row = workbook_sheet_row_dict(test_info_response.data, "Account Review")
        test_info_raw_headers = workbook_sheet_headers(test_info_response.data, "Raw Apify Data")
        test_info_raw_row = workbook_sheet_row_dict(test_info_response.data, "Raw Apify Data")
        assert {"identifier", "profile_url", "source_filename", "status", "reason"}.issubset(test_info_account_headers), test_info_account_headers
        assert {"identifier", "profile_url", "source_filename", "raw_data_source", "review_status", "review_reason"}.issubset(test_info_raw_headers), test_info_raw_headers
        assert test_info_account_row["identifier"] == "creatorgamma", test_info_account_row
        assert test_info_account_row["profile_url"] == "https://www.instagram.com/creatorgamma", test_info_account_row
        assert test_info_account_row["source_filename"] == "runtime_validation_upload.xlsx", test_info_account_row
        assert test_info_raw_row["identifier"] == "creatorgamma", test_info_raw_row
        assert test_info_raw_row["raw_data_source"] == "last_non_empty_snapshot", test_info_raw_row
        assert test_info_raw_row["source_filename"] == "runtime_validation_upload.xlsx", test_info_raw_row

        test_info_json_response = client.get("/api/download/instagram/test-info-json")
        assert test_info_json_response.status_code == 200, test_info_json_response.get_data(as_text=True)
        test_info_json_payload = json.loads(test_info_json_response.get_data(as_text=True))
        assert test_info_json_payload["raw_data"]["source"] == "last_non_empty_snapshot", test_info_json_payload
        profile_payload = next(item for item in test_info_json_payload["profiles"] if item["identifier"] == "creatorgamma")
        assert profile_payload["profile_url"] == "https://www.instagram.com/creatorgamma", profile_payload
        assert profile_payload["source_filename"] == "runtime_validation_upload.xlsx", profile_payload
        assert profile_payload["raw_data_source"] == "last_non_empty_snapshot", profile_payload
        assert profile_payload["status"] == "Pass", profile_payload

        saved_visual_artifact = backend_app.save_visual_review_artifacts(
            "instagram",
            {
                "visual_results": {
                    "creatorgamma": {
                        "username": "creatorgamma",
                        "decision": "Pass",
                        "reason": "saved visual pass",
                        "signals": ["clean"],
                        "success": True,
                    }
                },
                "review_history": [{"current_username": "creatorgamma", "step": "completed"}],
                "live_review": {"current_username": "creatorgamma", "step": "completed"},
            },
            metadata=backend_app.load_profile_review_artifact_metadata("instagram"),
        )
        assert saved_visual_artifact["visual_results_path"].endswith("instagram_visual_results.json"), saved_visual_artifact
        assert saved_visual_artifact["review_history_path"].endswith("instagram_visual_review_history.json"), saved_visual_artifact

        artifact_status_response = client.get("/api/artifacts/instagram/status")
        artifact_status_payload = artifact_status_response.get_json()
        assert artifact_status_response.status_code == 200, artifact_status_payload
        assert artifact_status_payload["saved_final_review_artifacts_available"] is True, artifact_status_payload
        assert artifact_status_payload["saved_final_review_artifacts_updated_at"], artifact_status_payload

        final_review_missing_payload = client.post("/api/download/instagram/final-review", json={})
        assert final_review_missing_payload.status_code == 200, final_review_missing_payload.get_data(as_text=True)
        final_review_missing_row = workbook_sheet_row_dict(final_review_missing_payload.data)
        assert final_review_missing_row["raw_data_source"] == "last_non_empty_snapshot", final_review_missing_row
        assert final_review_missing_row["stale_result"] is True, final_review_missing_row
        assert final_review_missing_row["visual_status"] == "Pass", final_review_missing_row

        final_review_empty_payload = client.post(
            "/api/download/instagram/final-review",
            json={"profile_reviews": [], "visual_results": {}},
        )
        assert final_review_empty_payload.status_code == 200, final_review_empty_payload.get_data(as_text=True)
        final_review_empty_row = workbook_sheet_row_dict(final_review_empty_payload.data)
        assert final_review_empty_row["raw_data_source"] == "last_non_empty_snapshot", final_review_empty_row
        assert final_review_empty_row["stale_result"] is True, final_review_empty_row

        final_review_incomplete_payload = client.post(
            "/api/download/instagram/final-review",
            json={
                "profile_reviews": [
                    {
                        "platform": "instagram",
                        "username": "creatorgamma",
                        "profile_url": "",
                        "status": "Pass",
                        "reason": "",
                        "covers": [],
                    }
                ],
                "visual_results": {},
            },
        )
        assert final_review_incomplete_payload.status_code == 200, final_review_incomplete_payload.get_data(as_text=True)
        final_review_incomplete_row = workbook_sheet_row_dict(final_review_incomplete_payload.data)
        assert final_review_incomplete_row["visual_status"] == "Pass", final_review_incomplete_row
        assert final_review_incomplete_row["reason"] == "saved visual pass", final_review_incomplete_row

        prescreen_review_response = client.get("/api/download/instagram/prescreen-review")
        assert prescreen_review_response.status_code == 200, prescreen_review_response.get_data(as_text=True)
        prescreen_headers = workbook_sheet_headers(prescreen_review_response.data)
        prescreen_row = workbook_sheet_row_dict(prescreen_review_response.data)
        assert {"identifier", "profile_url", "source_filename", "status", "reason"}.issubset(prescreen_headers), prescreen_headers
        assert {"stale_result", "raw_data_source"}.issubset(prescreen_headers), prescreen_headers
        assert prescreen_row["identifier"] == "creatorgamma", prescreen_row
        assert prescreen_row["source_filename"] == "runtime_validation_upload.xlsx", prescreen_row
        assert prescreen_row["stale_result"] is True, prescreen_row
        assert prescreen_row["raw_data_source"] == "last_non_empty_snapshot", prescreen_row

        image_review_response = client.get("/api/download/instagram/image-review")
        assert image_review_response.status_code == 200, image_review_response.get_data(as_text=True)
        image_review_headers = workbook_sheet_headers(image_review_response.data)
        image_review_row = workbook_sheet_row_dict(image_review_response.data)
        assert {"identifier", "profile_url", "source_filename", "status", "reason", "stage_status", "stage_reason"}.issubset(image_review_headers), image_review_headers
        assert {"stale_result", "raw_data_source"}.issubset(image_review_headers), image_review_headers
        assert image_review_row["identifier"] == "creatorgamma", image_review_row
        assert image_review_row["source_filename"] == "runtime_validation_upload.xlsx", image_review_row
        assert image_review_row["profile_url"] == "https://www.instagram.com/creatorgamma", image_review_row
        assert image_review_row["stale_result"] is True, image_review_row
        assert image_review_row["raw_data_source"] == "last_non_empty_snapshot", image_review_row

        downgraded_assets = backend_app.select_visual_review_bundle_assets(
            {
                "collages": [
                    {
                        "cover_count": 9,
                        "requested_cover_count": 9,
                        "download_failures": [],
                        "image_data_url": "data:image/jpeg;base64,collage-1",
                        "preview_bytes": b"collage-1",
                    },
                    {
                        "cover_count": 4,
                        "requested_cover_count": 9,
                        "download_failures": [],
                        "image_data_url": "data:image/jpeg;base64,collage-2",
                        "preview_bytes": b"collage-2",
                    },
                ],
                "cover_count": 13,
                "requested_cover_count": 18,
                "download_failures": [],
            },
            backend_app.VISUAL_REVIEW_MODE_AUTO,
        )
        assert downgraded_assets["requested_mode"] == backend_app.VISUAL_REVIEW_MODE_AUTO, downgraded_assets
        assert downgraded_assets["applied_mode"] == backend_app.VISUAL_REVIEW_MODE_SIMPLE, downgraded_assets
        assert downgraded_assets["collage_count"] == 2, downgraded_assets
        assert downgraded_assets["reviewed_collage_count"] == 1, downgraded_assets
        assert downgraded_assets["unused_collage_count"] == 1, downgraded_assets
        assert len(downgraded_assets["image_data_urls"]) == 1, downgraded_assets
        assert len(downgraded_assets["available_image_data_urls"]) == 2, downgraded_assets
        assert "自动降级" in downgraded_assets["downgrade_reason"], downgraded_assets

        def build_preview_url(job_id, username, _preview_bytes):
            return f"/api/visual-review-preview/{job_id}-{username}"

        insufficient_progress_events = []

        def fake_insufficient_bundle(_cover_urls, min_required_cover_count=1, max_collage_count=1, progress_callback=None):
            assert min_required_cover_count == backend_app.MIN_VISUAL_REVIEW_COVER_COUNT
            if progress_callback:
                progress_callback(
                    {
                        "type": "collage_error",
                        "collage_index": 1,
                        "collage_total": max_collage_count,
                        "aggregate_requested_cover_count": 9,
                        "aggregate_loaded_cover_count": 2,
                        "aggregate_failed_cover_count": 7,
                        "requested_cover_count": 9,
                        "current_cover_total": 9,
                        "current_cover_index": 9,
                        "min_required_cover_count": min_required_cover_count,
                        "error": "simulated collage failure",
                    }
                )
            return {
                "collages": [
                    {
                        "cover_count": 2,
                        "requested_cover_count": 9,
                        "download_failures": [],
                        "image_data_url": "data:image/jpeg;base64,scarce",
                        "preview_bytes": b"scarce",
                    }
                ],
                "collage_errors": [{"collage_index": 1, "error": "simulated collage failure"}],
                "cover_count": 2,
                "requested_cover_count": 9,
                "download_failure_count": 7,
                "download_failures": [{"host": "cdn.example.com", "error_type": "timeout", "error": "timeout"}],
                "stopped_early": True,
                "early_stop_reason": "剩余封面不足 3 张。",
                "min_required_cover_count": min_required_cover_count,
            }

        def fake_evaluate_should_not_run(*_args, **_kwargs):
            raise AssertionError("evaluate_cover_collage should not run when covers are insufficient")

        with patched_backend_attrs(
            build_cover_collage_bundle_assets=fake_insufficient_bundle,
            cache_visual_review_preview=build_preview_url,
            evaluate_cover_collage=fake_evaluate_should_not_run,
        ):
            insufficient_result = backend_app.perform_visual_review(
                "job-scarce",
                [{"username": "scarce", "covers": ["https://example.com/scarce-1.jpg"]}],
                platform="instagram",
                review_mode=backend_app.VISUAL_REVIEW_MODE_AUTO,
                progress_callback=lambda stage, message, **kwargs: insufficient_progress_events.append(
                    {"stage": stage, "message": message, **kwargs}
                ),
            )
        assert insufficient_result["summary"]["failed"] == 1, insufficient_result
        assert insufficient_result["review_history"][0]["step"] == "collage_failed", insufficient_result
        assert insufficient_result["review_history"][0]["collage_error_count"] == 1, insufficient_result
        assert insufficient_result["review_history"][0]["current_collage_urls"] == [
            "/api/visual-review-preview/job-scarce-scarce-collage-1"
        ], insufficient_result
        insufficient_live_review = insufficient_progress_events[-1]["partial_result"]["live_review"]
        assert insufficient_live_review["contract_version"] == "visual_review_live_v1", insufficient_live_review
        assert insufficient_live_review["step"] == "collage_failed", insufficient_live_review
        assert insufficient_live_review["requested_mode"] == backend_app.VISUAL_REVIEW_MODE_AUTO, insufficient_live_review
        assert insufficient_live_review["applied_mode"] == backend_app.VISUAL_REVIEW_MODE_SIMPLE, insufficient_live_review
        assert insufficient_live_review["reviewed_collage_count"] == 1, insufficient_live_review
        assert insufficient_live_review["collage_error_count"] == 1, insufficient_live_review
        assert insufficient_live_review["current_collage_urls"], insufficient_live_review

        multi_collage_progress_events = []

        def fake_success_bundle(_cover_urls, min_required_cover_count=1, max_collage_count=1, progress_callback=None):
            assert min_required_cover_count == backend_app.MIN_VISUAL_REVIEW_COVER_COUNT
            assert max_collage_count >= 2
            if progress_callback:
                progress_callback(
                    {
                        "type": "success",
                        "collage_index": 1,
                        "collage_total": 2,
                        "aggregate_requested_cover_count": 18,
                        "aggregate_loaded_cover_count": 18,
                        "aggregate_failed_cover_count": 0,
                        "requested_cover_count": 9,
                        "loaded_cover_count": 9,
                        "failed_cover_count": 0,
                        "current_cover_total": 9,
                        "current_cover_index": 9,
                        "cover_order_index": 9,
                        "min_required_cover_count": min_required_cover_count,
                    }
                )
            return {
                "collages": [
                    {
                        "cover_count": 9,
                        "requested_cover_count": 9,
                        "download_failures": [],
                        "image_data_url": "data:image/jpeg;base64,ok-1",
                        "preview_bytes": b"ok-1",
                    },
                    {
                        "cover_count": 9,
                        "requested_cover_count": 9,
                        "download_failures": [],
                        "image_data_url": "data:image/jpeg;base64,ok-2",
                        "preview_bytes": b"ok-2",
                    },
                ],
                "collage_errors": [],
                "cover_count": 18,
                "requested_cover_count": 18,
                "download_failure_count": 0,
                "download_failures": [],
                "stopped_early": False,
                "early_stop_reason": "",
                "min_required_cover_count": min_required_cover_count,
            }

        def fake_success_evaluate(username, cover_urls=None, collage_assets=None, platform="tiktok"):
            assert username == "rich"
            assert platform == "tiktok"
            assert cover_urls is None
            assert collage_assets["applied_mode"] == backend_app.VISUAL_REVIEW_MODE_ENHANCED, collage_assets
            assert len(collage_assets["image_data_urls"]) == 2, collage_assets
            return {
                "decision": "Pass",
                "reason": "多图风格稳定",
                "signals": ["多人互动", "产品展示"],
                "raw_text": "ok",
                "cover_count": 18,
                "collage_count": 2,
                "applied_mode": backend_app.VISUAL_REVIEW_MODE_ENHANCED,
                "downgrade_reason": "",
            }

        with patched_backend_attrs(
            build_cover_collage_bundle_assets=fake_success_bundle,
            cache_visual_review_preview=build_preview_url,
            evaluate_cover_collage=fake_success_evaluate,
        ):
            success_result = backend_app.perform_visual_review(
                "job-rich",
                [{"username": "rich", "covers": [f"https://example.com/rich-{idx}.jpg" for idx in range(1, 19)]}],
                platform="tiktok",
                review_mode=backend_app.VISUAL_REVIEW_MODE_ENHANCED,
                progress_callback=lambda stage, message, **kwargs: multi_collage_progress_events.append(
                    {"stage": stage, "message": message, **kwargs}
                ),
            )
        assert success_result["summary"]["passed"] == 1, success_result
        assert success_result["review_history"][0]["step"] == "completed", success_result
        assert success_result["review_history"][0]["collage_count"] == 2, success_result
        assert success_result["review_history"][0]["reviewed_collage_count"] == 2, success_result
        assert len(success_result["review_history"][0]["current_collage_urls"]) == 2, success_result
        success_live_review = multi_collage_progress_events[-1]["partial_result"]["live_review"]
        assert success_live_review["step"] == "completed", success_live_review
        assert success_live_review["reviewed_collage_count"] == 2, success_live_review
        assert success_live_review["current_collage_urls"] == [
            "/api/visual-review-preview/job-rich-rich-collage-1",
            "/api/visual-review-preview/job-rich-rich-collage-2",
        ], success_live_review
        assert success_live_review["signals"] == ["多人互动", "产品展示"], success_live_review

        os.environ["VISION_LEMONAPI_API_KEY"] = "lemonapi-test-key"
        os.environ["VISION_LEMONAPI_MODEL"] = "[L]gemini-3.1-pro-preview"
        provider_attempts = []
        lemonapi_base_url = "https://new.lemonapi.site/v1"
        lemonapi_model = "[L]gemini-3.1-pro-preview"
        test_collage_assets = {
            "image_data_urls": ["data:image/jpeg;base64,ok"],
            "cover_count": 9,
            "requested_mode": backend_app.VISUAL_REVIEW_MODE_AUTO,
            "applied_mode": backend_app.VISUAL_REVIEW_MODE_SIMPLE,
            "downgrade_reason": "",
        }

        class FakeChatCompletions:
            def __init__(self, base_url):
                self.base_url = base_url

            def create(self, model, messages):
                provider_attempts.append((self.base_url, model, messages))
                return {
                    "choices": [
                        {
                            "message": {
                                "content": '{"decision":"Pass","reason":"lemonapi ok","signals":["图像接口可用"]}'
                            }
                        }
                    ]
                }

        class FakeChat:
            def __init__(self, base_url):
                self.completions = FakeChatCompletions(base_url)

        class FakeResponsesNoop:
            def create(self, *args, **kwargs):
                raise AssertionError("responses API should not be used for lemonapi")

        class FakeOpenAIForLemonapi:
            def __init__(self, api_key, base_url, timeout):
                self.api_key = api_key
                self.base_url = base_url
                self.timeout = timeout
                self.chat = FakeChat(base_url)
                self.responses = FakeResponsesNoop()

        with patched_backend_attrs(OpenAI=FakeOpenAIForLemonapi):
            lemonapi_review = backend_app.evaluate_cover_collage(
                "lemonapi-demo",
                collage_assets=test_collage_assets,
                platform="tiktok",
            )
            assert lemonapi_review["provider"] == "lemonapi", lemonapi_review
            assert provider_attempts, provider_attempts
            requested_base_url, requested_model, requested_messages = provider_attempts[0]
            assert requested_base_url == lemonapi_base_url, provider_attempts
            assert requested_model == lemonapi_model, provider_attempts
            assert requested_messages[0]["content"][0]["type"] == "text", requested_messages
            assert requested_messages[0]["content"][1]["type"] == "image_url", requested_messages
            lemonapi_health = client.get("/api/health").get_json()
            assert lemonapi_health["checks"]["active_vision_provider"] == "lemonapi", lemonapi_health
            assert lemonapi_health["checks"]["vision_providers"] == ["lemonapi"], lemonapi_health

        os.environ.pop("VISION_LEMONAPI_API_KEY", None)
        os.environ.pop("VISION_LEMONAPI_MODEL", None)

        os.environ["VISION_QUAN2GO_API_KEY"] = "quan2go-test-key"
        os.environ["OPENAI_API_KEY"] = "auto-code-test-key"
        provider_attempts = []
        auto_code_base_url = "https://gpt.auto-code.net/rust/openai/v1"
        quan2go_base_url = "https://capi.quan2go.com/openai"
        test_collage_assets = {
            "image_data_urls": ["data:image/jpeg;base64,ok"],
            "cover_count": 9,
            "requested_mode": backend_app.VISUAL_REVIEW_MODE_AUTO,
            "applied_mode": backend_app.VISUAL_REVIEW_MODE_SIMPLE,
            "downgrade_reason": "",
        }

        class FakeResponses:
            def __init__(self, base_url):
                self.base_url = base_url

            def create(self, model, input):
                provider_attempts.append(self.base_url)
                if self.base_url == auto_code_base_url:
                    return '{"decision":"Pass","reason":"auto-code ok","signals":["备用接口可用"]}'
                raise RuntimeError(f"unexpected provider call: {self.base_url}")

        class FakeOpenAI:
            def __init__(self, api_key, base_url, timeout):
                self.api_key = api_key
                self.base_url = base_url
                self.timeout = timeout
                self.responses = FakeResponses(base_url)

        def fake_call_vision_stream(client, model, input_messages):
            provider_attempts.append(client.base_url)
            raise RuntimeError("quan2go unstable")

        with patched_backend_attrs(
            OpenAI=FakeOpenAI,
            call_vision_stream=fake_call_vision_stream,
        ):
            first_review = backend_app.evaluate_cover_collage(
                "fallback-demo",
                collage_assets=test_collage_assets,
                platform="tiktok",
            )
            assert first_review["provider"] == "auto-code", first_review
            assert provider_attempts == [quan2go_base_url, auto_code_base_url], provider_attempts
            assert backend_app.get_current_vision_provider_name() == "auto-code"
            first_health = client.get("/api/health").get_json()
            assert first_health["checks"]["active_vision_provider"] == "auto-code", first_health
            assert first_health["checks"]["vision_providers"] == ["auto-code", "quan2go"], first_health

            provider_attempts.clear()
            second_review = backend_app.evaluate_cover_collage(
                "fallback-demo-2",
                collage_assets=test_collage_assets,
                platform="tiktok",
            )
            assert second_review["provider"] == "auto-code", second_review
            assert provider_attempts == [auto_code_base_url], provider_attempts

        os.environ.pop("VISION_QUAN2GO_API_KEY", None)
        os.environ.pop("OPENAI_API_KEY", None)

        scrape_response = client.post("/api/jobs/scrape", json={"platform": "tiktok", "payload": {}})
        assert_json_error(scrape_response, "MISSING_APIFY_CONFIG", "APIFY_TOKEN")

        malformed_scrape = client.post("/api/jobs/scrape", json={"platform": "tiktok", "payload": []})
        malformed_payload = malformed_scrape.get_json()
        assert malformed_scrape.status_code == 400, malformed_payload
        assert malformed_payload["error"] == "payload 必须是对象", malformed_payload
        assert "error_code" not in malformed_payload, malformed_payload

        visual_response = client.post(
            "/api/jobs/visual-review",
            json={"profiles": [{"username": "demo", "covers": ["https://example.com/a.jpg"]}]},
        )
        assert_json_error(visual_response, "MISSING_VISION_CONFIG", "VISION_QUAN2GO_API_KEY")

        evaluate_response = client.post(
            "/api/evaluate_influencer",
            json={"username": "demo", "cover_urls": ["https://example.com/a.jpg"]},
        )
        assert_json_error(evaluate_response, "MISSING_VISION_CONFIG", "OPENAI_API_KEY")

    print("Runtime validation checks passed.")


if __name__ == "__main__":
    main()
